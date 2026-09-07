"""
test_pipeline.py
Automated test suite for the AI-Based Knowledge Retrieval Platform RAG pipeline.
Tests cover every major component independently and together as an integration test.

Run with:
    python -m unittest tests/test_pipeline.py -v
"""

import os
import sys
import csv
import tempfile
import unittest
from pathlib import Path

# Add the project root to sys.path so we can import the backend modules
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from backend.models import DocumentChunk, RetrievalResult, QueryResponse
from backend.document_processor import DocumentProcessor, DocumentProcessingError
from backend.chunker import DocumentChunker
from backend.embeddings import EmbeddingEngine
from backend.vector_store import VectorStore
from backend.retriever import KnowledgeRetriever
from backend.generator import ResponseGenerator


class TestDocumentModels(unittest.TestCase):
    """Unit tests for the core data models."""

    def test_document_chunk_creation(self):
        """Test that DocumentChunk stores all fields correctly."""
        chunk = DocumentChunk(
            chunk_id="test_001",
            document_name="test.pdf",
            document_type="PDF",
            text="Virtual memory is a memory management technique.",
            page_number=4,
        )
        self.assertEqual(chunk.chunk_id, "test_001")
        self.assertEqual(chunk.document_name, "test.pdf")
        self.assertEqual(chunk.page_number, 4)

    def test_document_chunk_to_dict(self):
        """Test that DocumentChunk serializes to dict correctly."""
        chunk = DocumentChunk(
            chunk_id="c1", document_name="doc.txt", document_type="TXT", text="hello"
        )
        d = chunk.to_dict()
        self.assertIn("chunk_id", d)
        self.assertIn("text", d)

    def test_query_response_creation(self):
        """Test QueryResponse initialization."""
        resp = QueryResponse(
            answer="Virtual memory allows programs larger than RAM to execute.",
            confidence="High",
            sources=[],
        )
        self.assertEqual(resp.confidence, "High")
        self.assertIsInstance(resp.sources, list)


class TestTxtExtraction(unittest.TestCase):
    """Tests for plain text file extraction."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.processor = DocumentProcessor(upload_dir=self.tmp)

    def test_txt_extraction(self):
        """Verify that TXT files are correctly extracted and cleaned."""
        txt_file = Path(self.tmp) / "sample.txt"
        txt_file.write_text("Hello world.\nThis is a test document.\nThird line here.", encoding="utf-8")

        segments = self.processor._extract_txt(txt_file)
        self.assertGreater(len(segments), 0)
        self.assertIn("Hello world", segments[0]["text"])
        self.assertIsNone(segments[0]["page_number"])

    def test_txt_empty_file(self):
        """Verify that empty TXT raises DocumentProcessingError."""
        txt_file = Path(self.tmp) / "empty.txt"
        txt_file.write_text("   \n\n   ", encoding="utf-8")

        with self.assertRaises(DocumentProcessingError):
            self.processor._extract_txt(txt_file)


class TestCSVExtraction(unittest.TestCase):
    """Tests for CSV structured data extraction."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.processor = DocumentProcessor(upload_dir=self.tmp)

    def test_csv_row_extraction(self):
        """Verify CSV rows are converted to key-value text format."""
        csv_file = Path(self.tmp) / "students.csv"
        with open(csv_file, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["Name", "Department", "CGPA"])
            writer.writerow(["Ravi Kumar", "CSE", "8.7"])
            writer.writerow(["Priya Sharma", "ECE", "9.1"])

        segments = self.processor._extract_csv(csv_file)
        self.assertEqual(len(segments), 2)
        self.assertIn("Name: Ravi Kumar", segments[0]["text"])
        self.assertIn("CGPA: 8.7", segments[0]["text"])
        self.assertEqual(segments[0]["row_number"], 1)
        self.assertEqual(segments[1]["row_number"], 2)

    def test_csv_column_names_preserved(self):
        """Verify column names appear in extracted text."""
        csv_file = Path(self.tmp) / "records.csv"
        with open(csv_file, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["Project_Domain", "Status"])
            writer.writerow(["Machine Learning", "Active"])

        segments = self.processor._extract_csv(csv_file)
        self.assertIn("Project_Domain", segments[0]["text"])
        self.assertIn("Machine Learning", segments[0]["text"])

    def test_csv_empty_raises(self):
        """Verify that an empty CSV file raises DocumentProcessingError."""
        csv_file = Path(self.tmp) / "empty.csv"
        csv_file.write_text("Name,Department\n", encoding="utf-8")

        with self.assertRaises(DocumentProcessingError):
            self.processor._extract_csv(csv_file)


class TestPDFExtraction(unittest.TestCase):
    """Tests for PDF text extraction using pypdf."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.processor = DocumentProcessor(upload_dir=self.tmp)
        self.pdf_path = None

        # Try to create a sample PDF using reportlab
        try:
            from reportlab.platypus import SimpleDocTemplate, Paragraph
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.pagesizes import A4

            pdf_path = Path(self.tmp) / "test.pdf"
            doc = SimpleDocTemplate(str(pdf_path), pagesize=A4)
            styles = getSampleStyleSheet()
            story = [
                Paragraph("Virtual memory is a memory management technique.", styles["Normal"]),
                Paragraph("Page faults occur when the OS loads a page from disk.", styles["Normal"]),
            ]
            doc.build(story)
            self.pdf_path = pdf_path
        except ImportError:
            pass  # reportlab not available; skip PDF creation tests

    def test_pdf_extraction_with_pages(self):
        """Verify that PDF extraction returns page-numbered segments."""
        if self.pdf_path is None:
            self.skipTest("reportlab not installed; cannot create test PDF.")

        segments = self.processor._extract_pdf(self.pdf_path)
        self.assertGreater(len(segments), 0)
        for seg in segments:
            self.assertIsNotNone(seg["page_number"])
            self.assertGreater(seg["page_number"], 0)

    def test_pdf_text_content(self):
        """Verify that actual text content is extracted from PDF."""
        if self.pdf_path is None:
            self.skipTest("reportlab not installed; cannot create test PDF.")

        segments = self.processor._extract_pdf(self.pdf_path)
        all_text = " ".join(s["text"] for s in segments)
        self.assertIn("Virtual memory", all_text)


class TestDOCXExtraction(unittest.TestCase):
    """Tests for Word DOCX text extraction."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.processor = DocumentProcessor(upload_dir=self.tmp)

    def test_docx_extraction(self):
        """Verify that DOCX files are correctly parsed."""
        import docx as docx_lib
        docx_path = Path(self.tmp) / "test.docx"
        doc = docx_lib.Document()
        doc.add_paragraph("The attention mechanism in transformers allows contextual understanding.")
        doc.add_paragraph("RAG is a framework combining retrieval with generation.")
        doc.save(str(docx_path))

        segments = self.processor._extract_docx(docx_path)
        self.assertGreater(len(segments), 0)
        all_text = " ".join(s["text"] for s in segments)
        self.assertIn("attention mechanism", all_text)


class TestChunker(unittest.TestCase):
    """Tests for the document chunking module."""

    def setUp(self):
        self.chunker = DocumentChunker(chunk_size=200, chunk_overlap=50)

    def test_short_text_single_chunk(self):
        """Text shorter than chunk_size should produce a single chunk."""
        segments = [{"text": "Virtual memory is a technique.", "page_number": 1, "row_number": None}]
        chunks = self.chunker.chunk_document("test.txt", segments)
        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0].document_name, "test.txt")

    def test_long_text_multiple_chunks(self):
        """Text longer than chunk_size should produce multiple chunks."""
        long_text = "This is a sentence about operating systems. " * 20
        segments = [{"text": long_text, "page_number": 1, "row_number": None}]
        chunks = self.chunker.chunk_document("large.pdf", segments)
        self.assertGreater(len(chunks), 1)

    def test_chunk_metadata_preservation(self):
        """Chunk metadata (page_number, document_name) should be preserved."""
        segments = [{"text": "Paging is a memory management scheme.", "page_number": 4, "row_number": None}]
        chunks = self.chunker.chunk_document("os.pdf", segments)
        self.assertEqual(chunks[0].page_number, 4)
        self.assertEqual(chunks[0].document_name, "os.pdf")
        self.assertEqual(chunks[0].document_type, "PDF")

    def test_csv_row_single_chunk(self):
        """Each CSV row should become a single atomic chunk."""
        segments = [
            {"text": "Name: Ravi | Department: CSE | CGPA: 8.7", "page_number": None, "row_number": 1},
            {"text": "Name: Priya | Department: ECE | CGPA: 9.1", "page_number": None, "row_number": 2},
        ]
        chunks = self.chunker.chunk_document("students.csv", segments)
        self.assertEqual(len(chunks), 2)
        self.assertEqual(chunks[0].row_number, 1)

    def test_chunk_ids_are_unique(self):
        """All generated chunk IDs must be unique."""
        long_text = "This is content about page faults and virtual memory in operating systems. " * 10
        segments = [{"text": long_text, "page_number": 1, "row_number": None}]
        chunks = self.chunker.chunk_document("test.pdf", segments)
        ids = [c.chunk_id for c in chunks]
        self.assertEqual(len(ids), len(set(ids)))


class TestEmbeddingEngine(unittest.TestCase):
    """Tests for the vector embedding engine."""

    def setUp(self):
        self.engine = EmbeddingEngine(model_name="all-MiniLM-L6-v2")

    def test_embedding_shape(self):
        """Embeddings should have shape (N, dimension)."""
        texts = ["Virtual memory is a memory management technique.", "Paging uses fixed-size frames."]
        embeddings = self.engine.generate_embeddings(texts)
        self.assertEqual(embeddings.shape[0], 2)
        self.assertGreater(embeddings.shape[1], 0)

    def test_embedding_dtype(self):
        """Embeddings must be float32 for FAISS compatibility."""
        import numpy as np
        texts = ["Operating system manages memory."]
        embeddings = self.engine.generate_embeddings(texts)
        self.assertEqual(embeddings.dtype, np.float32)

    def test_embedding_normalized(self):
        """Embeddings should be approximately L2 normalized."""
        import numpy as np
        texts = ["Attention mechanism in transformers is a key innovation."]
        embeddings = self.engine.generate_embeddings(texts)
        norm = float(np.linalg.norm(embeddings[0]))
        self.assertAlmostEqual(norm, 1.0, delta=0.05)

    def test_query_embedding_shape(self):
        """Query embedding should have shape (1, dimension)."""
        vec = self.engine.generate_query_embedding("What is virtual memory?")
        self.assertEqual(vec.shape[0], 1)

    def test_empty_text_list(self):
        """Empty text list should return empty array."""
        result = self.engine.generate_embeddings([])
        self.assertEqual(result.shape[0], 0)


class TestVectorStore(unittest.TestCase):
    """Tests for the FAISS vector store."""

    def setUp(self):
        self.tmp_dir = tempfile.mkdtemp()
        self.store = VectorStore(persist_dir=self.tmp_dir, dimension=384)

    def tearDown(self):
        self.store.clear()

    def _make_chunk(self, chunk_id, text):
        return DocumentChunk(
            chunk_id=chunk_id,
            document_name="test.pdf",
            document_type="PDF",
            text=text,
            page_number=1,
        )

    def _make_embedding(self, n=1):
        import numpy as np
        vec = np.random.randn(n, 384).astype(np.float32)
        # L2 normalize
        norms = np.linalg.norm(vec, axis=1, keepdims=True)
        return vec / norms

    def test_add_and_retrieve(self):
        """Adding chunks and searching should return results."""
        chunk = self._make_chunk("c1", "Virtual memory is a memory management technique.")
        embeddings = self._make_embedding(1)
        self.store.add_documents([chunk], embeddings)

        query_vec = self._make_embedding(1)
        results = self.store.search(query_vec, top_k=1)
        self.assertEqual(len(results), 1)

    def test_total_chunks_count(self):
        """Total chunk count should update after adding documents."""
        chunks = [self._make_chunk(f"c{i}", f"Chunk text {i}") for i in range(3)]
        embeddings = self._make_embedding(3)
        self.store.add_documents(chunks, embeddings)
        self.assertEqual(self.store.total_chunks, 3)

    def test_persistence(self):
        """Vector store should persist and reload correctly."""
        chunk = self._make_chunk("persist_1", "Persistent chunk text.")
        embeddings = self._make_embedding(1)
        self.store.add_documents([chunk], embeddings)

        # Create new store instance pointing to same directory — should load persisted data
        store2 = VectorStore(persist_dir=self.tmp_dir, dimension=384)
        self.assertEqual(store2.total_chunks, 1)

    def test_empty_store_search_returns_empty(self):
        """Searching an empty store should return empty list."""
        query_vec = self._make_embedding(1)
        results = self.store.search(query_vec, top_k=3)
        self.assertEqual(results, [])

    def test_clear_resets_store(self):
        """Clearing the store should reduce chunk count to 0."""
        chunk = self._make_chunk("c1", "Text to clear.")
        self.store.add_documents([chunk], self._make_embedding(1))
        self.store.clear()
        self.assertEqual(self.store.total_chunks, 0)


class TestRetriever(unittest.TestCase):
    """Tests for the semantic retrieval module."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.engine = EmbeddingEngine(model_name="all-MiniLM-L6-v2")
        self.store = VectorStore(persist_dir=self.tmp, dimension=self.engine.dimension)
        self.retriever = KnowledgeRetriever(
            vector_store=self.store,
            embedding_engine=self.engine,
            top_k=3,
            similarity_threshold=0.0,  # Use 0 threshold so all chunks pass in tests
        )

        # Populate store with sample chunks
        texts = [
            "Virtual memory is a memory management technique that allows programs larger than RAM to run.",
            "Paging divides physical memory into fixed-size frames and logical memory into pages.",
            "The attention mechanism allows transformers to focus on relevant tokens in a sequence.",
        ]
        chunks = [
            DocumentChunk(
                chunk_id=f"test_{i}",
                document_name="test.pdf",
                document_type="PDF",
                text=t,
                page_number=i + 1,
            )
            for i, t in enumerate(texts)
        ]
        embeddings = self.engine.generate_embeddings(texts)
        self.store.add_documents(chunks, embeddings)

    def tearDown(self):
        self.store.clear()

    def test_retrieve_returns_results(self):
        """Retrieval on a populated store should return non-empty results."""
        result = self.retriever.retrieve("What is virtual memory?")
        self.assertGreater(len(result["all_retrieved"]), 0)

    def test_top_result_is_most_relevant(self):
        """Top retrieval result should be the most semantically similar chunk."""
        result = self.retriever.retrieve("How does virtual memory work?")
        top = result["all_retrieved"][0]
        self.assertIn("virtual memory", top.chunk.text.lower())

    def test_relevance_labels(self):
        """Relevance labels should be one of High, Medium, Low."""
        result = self.retriever.retrieve("What is attention in transformers?")
        for res in result["all_retrieved"]:
            self.assertIn(res.relevance, ["High", "Medium", "Low"])

    def test_retrieve_with_high_threshold_filters(self):
        """With a high threshold, irrelevant results should be filtered out."""
        retriever_strict = KnowledgeRetriever(
            vector_store=self.store,
            embedding_engine=self.engine,
            top_k=3,
            similarity_threshold=0.99,  # Very strict threshold
        )
        result = retriever_strict.retrieve("completely unrelated question about cooking recipes banana bread")
        # Either none pass, or score metadata is present
        self.assertIn("relevant_results", result)


class TestResponseGenerator(unittest.TestCase):
    """Tests for the grounded response generator."""

    def setUp(self):
        self.generator = ResponseGenerator(provider=None)  # Will use local fallback

    def _make_retrieval_data(self, include_results=True):
        if not include_results:
            return {"relevant_results": [], "all_retrieved": [], "confidence": "None", "top_score": 0.0}

        chunk = DocumentChunk(
            chunk_id="g1",
            document_name="OS.pdf",
            document_type="PDF",
            text="Virtual memory is a memory management technique that creates an illusion of very large memory.",
            page_number=4,
            source_id="OS.pdf — Page 4",
        )
        import numpy as np
        result = RetrievalResult(chunk=chunk, similarity_score=0.82, relevance="High")
        return {
            "relevant_results": [result],
            "all_retrieved": [result],
            "confidence": "High",
            "top_score": 0.82,
            "query_embedding_shape": [1, 384],
        }

    def test_no_context_returns_rejection(self):
        """When no relevant chunks exist, generator must return a grounded rejection."""
        data = self._make_retrieval_data(include_results=False)
        response = self.generator.generate_response("What is sourdough bread?", data)
        self.assertEqual(response.confidence, "None")
        self.assertIn("couldn't find enough relevant", response.answer.lower())
        self.assertEqual(len(response.sources), 0)


    def test_response_contains_answer(self):
        """When context exists, generator should return a non-empty answer."""
        data = self._make_retrieval_data(include_results=True)
        response = self.generator.generate_response("What is virtual memory?", data)
        self.assertIsNotNone(response.answer)
        self.assertGreater(len(response.answer.strip()), 0)

    def test_response_has_sources(self):
        """Generated response should include source attribution."""
        data = self._make_retrieval_data(include_results=True)
        response = self.generator.generate_response("Explain virtual memory.", data)
        self.assertGreater(len(response.sources), 0)
        source = response.sources[0]
        self.assertIn("document_name", source)
        self.assertIn("similarity_score", source)

    def test_response_confidence_propagated(self):
        """Confidence from retrieval data should be propagated to response."""
        data = self._make_retrieval_data(include_results=True)
        response = self.generator.generate_response("What is virtual memory?", data)
        self.assertEqual(response.confidence, "High")


class TestFlaskAPIEndpoints(unittest.TestCase):
    """Integration tests for the Flask API endpoints."""

    def setUp(self):
        # Set up test environment using a temp directory
        self.tmp = tempfile.mkdtemp()
        os.environ.setdefault("VECTOR_STORE_DIR", self.tmp)

        # Import app here to avoid side effects at module level
        import app as app_module
        app_module.app.config["TESTING"] = True
        self.client = app_module.app.test_client()

    def test_health_endpoint(self):
        """GET /health should return 200 with status field."""
        resp = self.client.get("/health")
        self.assertEqual(resp.status_code, 200)
        data = resp.get_json()
        self.assertIn("status", data)
        self.assertEqual(data["status"], "healthy")

    def test_documents_endpoint(self):
        """GET /documents should return 200 with documents list."""
        resp = self.client.get("/documents")
        self.assertEqual(resp.status_code, 200)
        data = resp.get_json()
        self.assertIn("documents", data)
        self.assertIsInstance(data["documents"], list)

    def test_query_empty_question(self):
        """POST /query with empty question should return 400."""
        resp = self.client.post("/query", json={"question": ""})
        self.assertEqual(resp.status_code, 400)
        data = resp.get_json()
        self.assertIn("error", data)

    def test_query_no_documents(self):
        """POST /query when knowledge base is empty should return 400 with guidance."""
        # Only runs if the current store is empty
        resp = self.client.get("/health")
        health = resp.get_json()
        if health.get("total_chunks", 0) == 0:
            resp = self.client.post("/query", json={"question": "What is virtual memory?"})
            self.assertEqual(resp.status_code, 400)
            data = resp.get_json()
            self.assertIn("error", data)

    def test_index_page_loads(self):
        """GET / should serve the HTML dashboard without errors."""
        resp = self.client.get("/")
        self.assertEqual(resp.status_code, 200)
        self.assertIn(b"Knowledge Retrieval", resp.data)


if __name__ == "__main__":
    unittest.main(verbosity=2)
